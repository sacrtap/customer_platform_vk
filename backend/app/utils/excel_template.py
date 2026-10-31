"""Excel 导入模板的公共生成器"""

import io

from sanic.response import raw

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def template_response(
    columns: list[tuple[str, str]],
    filename: str,
    sheet_name: str = "导入模板",
):
    """生成导入模板并返回 raw 响应（GET /import-template 端点统一出口）。

    模板结构（与 read_import_dataframe 的说明行剥离逻辑配套）：
        第 1 行英文列名（与解析列名一致）
        第 2 行中文说明，必须以「必填：」/「可选：」开头

    **不写示例数据行**：read_import_dataframe 只丢弃第 2 行中文说明行，第 3 行
    示例数据会被当作真实数据导入。用户下载模板后通常直接在示例行下方续写，
    示例行会被静默创建成一条记录（引用不存在的客户编号时还会整行报错，
    干扰用户判断）。示例值请并入第 2 行说明文字（如「可选：备注，如季度返点」）。

    Args:
        columns: [(英文字段名, 中文说明), ...]，列序即 Excel 列序
        filename: 下载文件名（如 "行业类型导入模板.xlsx"）
        sheet_name: 工作表名，默认「导入模板」

    注意：第 2 行的「必填：」/「可选：」前缀是 `_is_template_note_row` 判定并
    丢弃说明行的依据，新增模板时不可省略或改写成其他前缀。
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    assert ws is not None  # 新建 Workbook 必有活动工作表
    ws.title = sheet_name
    for i, (en_name, cn_note) in enumerate(columns, start=1):
        ws.cell(row=1, column=i, value=en_name)
        ws.cell(row=2, column=i, value=cn_note)
        ws.column_dimensions[get_column_letter(i)].width = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return raw(
        output.read(),
        content_type=XLSX_MIME,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
