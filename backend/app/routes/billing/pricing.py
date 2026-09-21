"""定价规则路由 — CRUD 和冲突检测"""

import asyncio
import io
import json as _json
import logging
from datetime import datetime
from typing import Any

from sanic.request import Request
from sanic.response import json, raw
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from ...cache.base import cache_service
from ...middleware.auth import auth_required, get_current_user, require_permission
from ...repository import PricingRepository
from ...services.billing import PricingService
from ...utils.audit_helpers import build_batch_audit_summary, create_audit_entry
from ...utils.excel_import import read_import_dataframe
from ...utils.tiers import parse_tiers_or_raise
from ...utils.timezone import local_date_to_utc_end, local_date_to_utc_start, utc_to_cst_date_str
from . import billing_bp

logger = logging.getLogger(__name__)


@billing_bp.get("/pricing-rules")
@auth_required
@require_permission("billing:view")
async def get_pricing_rules(request: Request):
    """获取定价规则列表（支持分页）"""
    db: AsyncSession = request.ctx.db_session
    pricing_service = PricingService(PricingRepository(db))

    # 分页参数
    page = int(request.args.get("page", 1))
    page_size = int(request.args.get("page_size", 20))
    page_size = min(page_size, 100)

    # 筛选参数
    try:
        customer_id = (
            int(request.args.get("customer_id")) if request.args.get("customer_id") else None
        )
    except ValueError:
        return json({"code": 40001, "message": "customer_id 参数必须为整数"}, status=400)
    keyword = request.args.get("keyword")  # 客户名称模糊搜索
    device_type = request.args.get("device_type")
    layer_type = request.args.get("layer_type")
    pricing_type = request.args.get("pricing_type")

    rules, total = await pricing_service.get_pricing_rules(
        customer_id=customer_id,
        keyword=keyword,
        device_type=device_type,
        layer_type=layer_type,
        pricing_type=pricing_type,
        page=page,
        page_size=page_size,
    )

    return json(
        {
            "code": 0,
            "message": "success",
            "data": {
                "list": [
                    {
                        "id": r.id,
                        "customer_id": r.customer_id,
                        "customer_name": r.customer.name if r.customer else None,
                        "device_type": r.device_type,
                        "layer_type": r.layer_type or "single",
                        "pricing_type": r.pricing_type,
                        # 真值判断会把 0 元单价/加层单价当成 None，须显式判 None
                        "unit_price": float(r.unit_price) if r.unit_price is not None else None,  # pyright: ignore[reportArgumentType, reportGeneralTypeIssues]
                        "multi_floor_pricing_type": r.multi_floor_pricing_type,  # pyright: ignore[reportAttributeAccessIssue]
                        "additional_floor_price": (
                            float(r.additional_floor_price)
                            if r.additional_floor_price is not None
                            else None  # pyright: ignore[reportAttributeAccessIssue, reportArgumentType, reportGeneralTypeIssues]
                        ),
                        "tiers": r.tiers,
                        "package_type": r.package_type,
                        "package_limits": r.package_limits,
                        "effective_date": (
                            r.effective_date.isoformat() if r.effective_date else None  # pyright: ignore[reportGeneralTypeIssues]
                        ),
                        "expiry_date": r.expiry_date.isoformat() if r.expiry_date else None,  # pyright: ignore[reportGeneralTypeIssues]
                    }
                    for r in rules
                ],
                "total": total,
                "page": page,
                "page_size": page_size,
            },
        }
    )


@billing_bp.post("/pricing-rules")
@auth_required
@require_permission("billing:edit")
async def create_pricing_rule(request: Request):
    """
    创建定价规则

    Body:
    {
        "customer_id": 1,
        "device_type": "X",
        "pricing_type": "fixed",
        "unit_price": 10.00,
        "effective_date": "2026-04-01",
        "expiry_date": "2026-12-31"
    }
    """
    db: AsyncSession = request.ctx.db_session
    data = request.json
    user = get_current_user(request)

    pricing_service = PricingService(PricingRepository(db))
    data["created_by"] = user["user_id"] if user else 1

    # 日期转换：前端本地日期 → UTC datetime
    if "effective_date" in data and isinstance(data["effective_date"], str):
        data["effective_date"] = local_date_to_utc_start(data["effective_date"])
    if "expiry_date" in data and isinstance(data["expiry_date"], str):
        data["expiry_date"] = local_date_to_utc_end(data["expiry_date"])

    try:
        rule = await pricing_service.create_pricing_rule(data)
    except ValueError as e:
        return json({"code": 40001, "message": str(e)}, status=400)

    # 定价规则变更后清除相关缓存
    await cache_service.invalidate_billing_cache()

    return json(
        {
            "code": 0,
            "message": "创建成功",
            "data": {
                "id": rule.id,
                "device_type": rule.device_type,
                "layer_type": rule.layer_type or "single",
                "pricing_type": rule.pricing_type,
                # 真值判断会把 0 元单价/加层单价当成 None，须显式判 None
                "unit_price": float(rule.unit_price) if rule.unit_price is not None else None,  # pyright: ignore[reportArgumentType, reportGeneralTypeIssues]
                "multi_floor_pricing_type": rule.multi_floor_pricing_type,  # pyright: ignore[reportAttributeAccessIssue]
                "additional_floor_price": (
                    float(rule.additional_floor_price)
                    if rule.additional_floor_price is not None
                    else None  # pyright: ignore[reportAttributeAccessIssue, reportArgumentType, reportGeneralTypeIssues]
                ),
            },
        },
        status=201,
    )


@billing_bp.put("/pricing-rules/<rule_id:int>")
@auth_required
@require_permission("billing:edit")
async def update_pricing_rule(request: Request, rule_id: int):
    """更新定价规则"""
    db: AsyncSession = request.ctx.db_session
    data = request.json

    pricing_service = PricingService(PricingRepository(db))

    # 日期转换：前端本地日期 → UTC datetime
    if "effective_date" in data and isinstance(data["effective_date"], str):
        data["effective_date"] = local_date_to_utc_start(data["effective_date"])
    if "expiry_date" in data and isinstance(data["expiry_date"], str):
        data["expiry_date"] = local_date_to_utc_end(data["expiry_date"])

    try:
        rule = await pricing_service.update_pricing_rule(rule_id, data)
    except ValueError as e:
        return json({"code": 40001, "message": str(e)}, status=400)

    if not rule:
        return json({"code": 40401, "message": "规则不存在"}, status=404)

    # 定价规则变更后清除相关缓存
    await cache_service.invalidate_billing_cache()

    return json(
        {
            "code": 0,
            "message": "更新成功",
            "data": {
                "id": rule.id,
                "device_type": rule.device_type,
                "layer_type": rule.layer_type or "single",
                "pricing_type": rule.pricing_type,
            },
        }
    )


@billing_bp.delete("/pricing-rules/<rule_id:int>")
@auth_required
@require_permission("billing:delete")
async def delete_pricing_rule(request: Request, rule_id: int):
    """删除定价规则"""
    db: AsyncSession = request.ctx.db_session
    pricing_service = PricingService(PricingRepository(db))

    success = await pricing_service.delete_pricing_rule(rule_id)

    if not success:
        return json({"code": 40401, "message": "规则不存在"}, status=404)

    # 定价规则变更后清除相关缓存
    await cache_service.invalidate_billing_cache()

    return json({"code": 0, "message": "删除成功"})


@billing_bp.get("/pricing-rules/check-conflict")
@auth_required
@require_permission("billing:view")
async def check_pricing_rule_conflict(request: Request):
    """
    检查定价规则有效期冲突

    Query params:
    - customer_id (必填, int)
    - pricing_type (必填, string) — fixed/tiered/package
    - device_type (可选, string) — 包年结算时可为空
    - layer_type (可选, string)
    - effective_date (必填, date)
    - expiry_date (可选, date)
    - exclude_id (可选, int) — 编辑时排除自身
    """
    db: AsyncSession = request.ctx.db_session

    # 参数校验
    try:
        customer_id = int(request.args.get("customer_id", 0))
        pricing_type = request.args.get("pricing_type")
        device_type = request.args.get("device_type")  # 可选，包年结算时可为 None
        layer_type = request.args.get("layer_type")
        effective_date_str = request.args.get("effective_date", "")
        expiry_date_str = request.args.get("expiry_date")
        exclude_id_str = request.args.get("exclude_id")

        if not customer_id or not pricing_type or not effective_date_str:
            return json(
                {
                    "code": 40001,
                    "message": "缺少必填参数：customer_id, pricing_type, effective_date",
                },
                status=400,
            )

        effective_date = local_date_to_utc_start(effective_date_str)
        expiry_date = local_date_to_utc_end(expiry_date_str) if expiry_date_str else None
        exclude_id = int(exclude_id_str) if exclude_id_str else None
    except (ValueError, TypeError):
        return json(
            {
                "code": 40001,
                "message": "参数格式错误",
            },
            status=400,
        )

    pricing_service = PricingService(PricingRepository(db))

    conflicting_rules = await pricing_service.check_pricing_rule_conflict(
        customer_id=customer_id,
        pricing_type=pricing_type,
        device_type=device_type,
        layer_type=layer_type,
        effective_date=effective_date,
        expiry_date=expiry_date,
        exclude_id=exclude_id,
    )

    return json(
        {
            "code": 0,
            "data": {
                "has_conflict": len(conflicting_rules) > 0,
                "conflicting_rules": [
                    {
                        "id": r.id,
                        "pricing_type": r.pricing_type,
                        "effective_date": (
                            r.effective_date.isoformat() if r.effective_date else None  # pyright: ignore[reportGeneralTypeIssues]
                        ),
                        "expiry_date": r.expiry_date.isoformat() if r.expiry_date else None,  # pyright: ignore[reportGeneralTypeIssues]
                    }
                    for r in conflicting_rules
                ],
            },
        }
    )


# ==================== 计费规则导入导出 ====================


@billing_bp.post("/pricing-rules/import")
@auth_required
@require_permission("billing:pricing_import")
async def import_pricing_rules(request: Request):
    """
    Excel 批量导入计费规则

    Form:
    - file: Excel 文件 (.xlsx)

    Excel 列要求:
    - company_id (必填) - 客户编号
    - pricing_type (必填) - fixed/tiered/package
    - effective_date (必填) - 生效日期 YYYY-MM-DD
    - device_type (非包年必填) - X/N/L（包年结算可为空）
    - layer_type (非包年必填) - single/multi/single_and_multi
    - unit_price (可选) - 定价单价
    - additional_floor_price (可选) - 加层单价
    - multi_floor_pricing_type (可选) - unified/incremental
    - tiers (可选) - 阶梯配置 JSON 字符串，如 [{"min":0,"max":null,"price":5}]
    - package_type (可选) - A/B/C/D（包年结算必填）
    - expiry_date (可选) - 失效日期 YYYY-MM-DD
    """
    import pandas as pd

    from ...models.billing import PackagePlan
    from ...models.customers import Customer

    files = request.files
    if "file" not in files:  # pyright: ignore[reportOperatorIssue]
        return json({"code": 40001, "message": "请上传 Excel 文件"}, status=400)

    excel_file = files["file"][0]  # pyright: ignore[reportOptionalSubscript]
    if not excel_file.name.endswith(".xlsx"):
        return json({"code": 40002, "message": "请上传 .xlsx 格式的文件"}, status=400)

    try:
        # 读取 Excel 文件（自动丢弃模板第 2 行的中文说明行）。pd.read_excel 为 CPU+I/O
        # 密集操作，10MB xlsx 解析可能阻塞事件循环数百毫秒到数秒，移到线程执行。
        df = await asyncio.to_thread(read_import_dataframe, excel_file.body, "company_id")

        # 必填列检查
        required_columns = ["company_id", "pricing_type", "effective_date"]
        missing_columns = [col for col in required_columns if col not in df.columns]
        if missing_columns:
            return json(
                {
                    "code": 40003,
                    "message": f"Excel 缺少必填列：{', '.join(missing_columns)}",
                },
                status=400,
            )

        # 行数限制
        if len(df) > 1000:
            return json(
                {"code": 40004, "message": "单次最多导入 1000 条记录"},
                status=400,
            )

        db_session: AsyncSession = request.ctx.db_session

        # 预加载所有客户 company_id -> customer_id 映射（排除软删除客户：
        # 软删除客户已不可在页面选择，导入若仍可命中会为其创建规则，导出后也无法回灌）
        result = await db_session.execute(
            select(Customer.id, Customer.company_id).where(Customer.deleted_at.is_(None))
        )
        company_to_customer = {row[1]: row[0] for row in result.all()}

        # 预加载有效的包年套餐类型（仅 active 且未软删除，与 create_pricing_rule 查询条件一致）
        plan_result = await db_session.execute(
            select(PackagePlan.package_type).where(
                PackagePlan.deleted_at.is_(None),
                PackagePlan.status == "active",
            )
        )
        active_package_types = set(plan_result.scalars().all())

        current_user = get_current_user(request)
        operator_id = current_user.get("user_id") if current_user else 1

        pricing_service = PricingService(PricingRepository(db_session))

        # 逐行校验并创建
        errors = []
        success_count = 0
        for idx, row in df.iterrows():
            # pandas-stubs 将 Series.get 的返回值推断为 Dtype，导致 pd.isna(...)
            # 误报 reportCallIssue/reportArgumentType；运行时是真实单元格值。
            # 标注 Any 消除误报，不改变语义。
            row: Any = row
            row_num = int(str(idx)) + 2  # Excel 行号（含表头）
            try:
                # 校验 company_id
                company_id = row.get("company_id")
                if bool(pd.isna(company_id)) or company_id is None:
                    errors.append(f"第 {row_num} 行：客户编号为空")
                    continue
                try:
                    # 先拒绝非整数值：Excel 单元格写 100001.9 会被 int() 静默截断成
                    # 100001，命中错误客户并静默创建规则；字符串 "100001.0" 等旧行为不变。
                    if isinstance(company_id, float) and not company_id.is_integer():
                        raise ValueError
                    company_id = int(company_id)
                except (ValueError, TypeError):
                    errors.append(f"第 {row_num} 行：客户编号 '{company_id}' 不是有效整数")
                    continue
                if company_id not in company_to_customer:
                    errors.append(f"第 {row_num} 行：客户编号 {company_id} 不存在")
                    continue

                # 校验 pricing_type
                pricing_type = str(row.get("pricing_type", "")).strip().lower()
                if pricing_type not in ("fixed", "tiered", "package"):
                    errors.append(f"第 {row_num} 行：计费类型必须为 fixed/tiered/package")
                    continue

                # 校验 effective_date
                effective_date_raw = row.get("effective_date")
                if bool(pd.isna(effective_date_raw)) or effective_date_raw is None:
                    errors.append(f"第 {row_num} 行：生效日期为空")
                    continue
                try:
                    effective_date = local_date_to_utc_start(str(effective_date_raw)[:10])
                except (ValueError, TypeError):
                    errors.append(f"第 {row_num} 行：生效日期格式错误")
                    continue

                # 可选字段
                device_type_raw = row.get("device_type")
                device_type = (
                    str(device_type_raw).strip() if not bool(pd.isna(device_type_raw)) else None
                )
                layer_type_raw = row.get("layer_type")
                layer_type = (
                    str(layer_type_raw).strip() if not bool(pd.isna(layer_type_raw)) else None
                )
                unit_price_raw = row.get("unit_price")
                unit_price = (
                    float(unit_price_raw)
                    if unit_price_raw is not None and not bool(pd.isna(unit_price_raw))
                    else None
                )
                additional_floor_price_raw = row.get("additional_floor_price")
                additional_floor_price = (
                    float(additional_floor_price_raw)
                    if additional_floor_price_raw is not None
                    and not bool(pd.isna(additional_floor_price_raw))
                    else None
                )
                multi_floor_pricing_type_raw = row.get("multi_floor_pricing_type")
                multi_floor_pricing_type = (
                    str(multi_floor_pricing_type_raw).strip()
                    if not bool(pd.isna(multi_floor_pricing_type_raw))
                    else None
                )
                # 多楼层计费方式：只允许 unified/incremental（与模型字段/UI 下拉/模板说明一致）。
                # 非法值此前会被原样落库，结算时按「非 incremental 即 unified」静默处理，
                # 用户以为配置生效但金额口径不是所选方式。
                if multi_floor_pricing_type and multi_floor_pricing_type not in (
                    "unified",
                    "incremental",
                ):
                    errors.append(f"第 {row_num} 行：多楼层计费方式必须为 unified/incremental")
                    continue
                package_type_raw = row.get("package_type")
                package_type = (
                    str(package_type_raw).strip() if not bool(pd.isna(package_type_raw)) else None
                )

                # tiers JSON 解析（归一化与校验由 parse_tiers_or_raise 统一处理）
                tiers = None
                tiers_raw = row.get("tiers")
                if tiers_raw is not None and not bool(pd.isna(tiers_raw)):
                    try:
                        parsed = _json.loads(tiers_raw) if isinstance(tiers_raw, str) else tiers_raw
                        tiers = parse_tiers_or_raise(parsed, row_num=row_num)
                    except _json.JSONDecodeError as e:
                        # JSON 本身不合法（JSONDecodeError 是 ValueError 子类，须先捕获）
                        errors.append(f"第 {row_num} 行：阶梯配置 JSON 格式错误：{e}")
                        continue
                    except ValueError as e:
                        # parse_tiers_or_raise 已生成「第 N 行：阶梯配置…」行级文案
                        errors.append(str(e))
                        continue

                # expiry_date
                expiry_date = None
                expiry_date_raw = row.get("expiry_date")
                if expiry_date_raw is not None and not bool(pd.isna(expiry_date_raw)):
                    try:
                        expiry_date = local_date_to_utc_end(str(expiry_date_raw)[:10])
                    except (ValueError, TypeError):
                        errors.append(f"第 {row_num} 行：失效日期格式错误")
                        continue

                # 包年结算必填 package_type
                if pricing_type == "package" and not package_type:
                    errors.append(f"第 {row_num} 行：包年结算必须填写套餐类型")
                    continue

                # 包年结算：套餐类型必须存在且 status='active'（否则会静默创建 unit_price=None 的规则）
                if pricing_type == "package" and package_type not in active_package_types:
                    errors.append(f"第 {row_num} 行：套餐类型 '{package_type}' 不存在或已停用")
                    continue

                # 非包年结算：设备类型与楼层类型必填且取值合法（与 UI 表单 required + 下拉选项一致）
                if pricing_type != "package":
                    if not device_type:
                        errors.append(f"第 {row_num} 行：设备类型不能为空（非包年结算必填）")
                        continue
                    if device_type not in ("X", "N", "L"):
                        errors.append(f"第 {row_num} 行：设备类型必须为 X/N/L")
                        continue
                    if not layer_type:
                        errors.append(f"第 {row_num} 行：楼层类型不能为空（非包年结算必填）")
                        continue
                    if layer_type not in ("single", "multi", "single_and_multi"):
                        errors.append(
                            f"第 {row_num} 行：楼层类型必须为 single/multi/single_and_multi"
                        )
                        continue

                # 单价不允许为负：负单价会直接参与结算金额计算（产生负向账单），
                # UI 表单输入框同样限制非负，导入端不应放宽。
                if unit_price is not None and unit_price < 0:
                    errors.append(f"第 {row_num} 行：单价不能为负数")
                    continue
                if additional_floor_price is not None and additional_floor_price < 0:
                    errors.append(f"第 {row_num} 行：加层单价不能为负数")
                    continue

                # 定价内容完整性：与 UI 表单对齐（fixed 必填 unit_price，tiered 必填至少一条阶梯），
                # 否则会静默按 0 元结算
                if pricing_type == "fixed" and unit_price is None:
                    errors.append(f"第 {row_num} 行：定价结算必须填写单价")
                    continue
                if pricing_type == "tiered" and not tiers:
                    errors.append(f"第 {row_num} 行：阶梯结算必须至少配置一条阶梯")
                    continue

                rule_data = {
                    "customer_id": company_to_customer[company_id],
                    "pricing_type": pricing_type,
                    "effective_date": effective_date,
                    "expiry_date": expiry_date,
                    "device_type": device_type,
                    "layer_type": layer_type,
                    "unit_price": unit_price,
                    "additional_floor_price": additional_floor_price,
                    "multi_floor_pricing_type": multi_floor_pricing_type,
                    "tiers": tiers,
                    "package_type": package_type,
                    "created_by": operator_id,
                }

                await pricing_service.create_pricing_rule(rule_data)
                success_count += 1
            except ValueError as e:
                # create_pricing_rule 内部逐行 commit，失败行的工作本就未落库；
                # 但 DB 级异常会让会话进入 PendingRollback，必须回滚才能继续后续行，
                # 否则后续所有行与最终审计 commit 全部失败 → 整批报 500 而前面的行已永久落库。
                # 此处不能用 begin_nested()：create_pricing_rule 内部的 commit 会直接释放保存点。
                await db_session.rollback()
                errors.append(f"第 {row_num} 行：{str(e)}")
            except Exception as e:  # 兜底，避免单行异常中断整个导入
                await db_session.rollback()
                # 通用异常只记录日志，不回传 str(e)：DB 约束名/驱动报错可能泄露内部信息到前端
                logger.warning("计费规则导入第 %d 行失败: %s", row_num, e)
                errors.append(f"第 {row_num} 行：导入失败，请检查该行数据")

        # 清除计费规则相关缓存
        await cache_service.invalidate_billing_cache()

        # 记录审计日志
        summary = build_batch_audit_summary(
            operation="pricing_rule_import",
            total_count=len(df),
            success_count=success_count,
            failed_count=len(errors),
            details=errors[:10],
        )

        await create_audit_entry(
            db_session=db_session,
            user_id=operator_id,
            action="batch_create",
            module="billing",
            record_id=None,
            record_type="pricing_rule",
            changes={"after": {"count": success_count}},
            operation_type="batch",
            extra_metadata=summary,
            ip_address=request.headers.get(
                "x-real-ip", request.headers.get("x-forwarded-for", request.ip)
            ),
            auto_commit=True,
        )

        return json(
            {
                "code": 0,
                "message": "导入完成",
                "data": {
                    "success_count": success_count,
                    "error_count": len(errors),
                    "errors": errors[:10],
                },
            }
        )
    except Exception as e:
        # 外层失败（如审计 commit 抛错）也需回滚，避免把中毒会话归还连接池
        await request.ctx.db_session.rollback()
        return json({"code": 50001, "message": f"导入失败：{str(e)}"}, status=500)


@billing_bp.get("/pricing-rules/import-template")
@auth_required
async def download_pricing_rule_import_template(request: Request):
    """下载计费规则导入 Excel 模板"""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    assert ws is not None  # 新建 Workbook 必有活动工作表
    ws.title = "计费规则导入模板"

    headers = [
        "company_id",
        "pricing_type",
        "effective_date",
        "device_type",
        "layer_type",
        "unit_price",
        "additional_floor_price",
        "multi_floor_pricing_type",
        "tiers",
        "package_type",
        "expiry_date",
    ]
    ws.append(headers)  # pyright: ignore[reportOptionalMemberAccess]

    notes = [
        "必填：客户编号（整数）",
        "必填：fixed/tiered/package",
        "必填：YYYY-MM-DD",
        "必填：X/N/L（非包年必填，包年可空）",
        "必填：single/multi/single_and_multi（非包年必填）",
        "可选：单价",
        "可选：加层单价",
        "可选：unified/incremental",
        '可选：JSON 数组，如 [{"min":0,"max":null,"price":5}]（首档 min 须为 0）',
        "可选：A/B/C/D（包年必填）",
        "可选：YYYY-MM-DD",
    ]
    ws.append(notes)  # pyright: ignore[reportOptionalMemberAccess]

    for col_num in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_num)].width = 26

    # 不再追加示例数据行：read_import_dataframe 只丢弃第 2 行（说明行），第 3 行会被当作
    # 真实数据导入。用户下载模板后通常直接在示例行下方续写，示例行会被静默创建成一条
    # 计费规则（客户编号 100001 不存在时还会整行报错，干扰用户判断）。
    # 如需示例，表头 + 说明行已足以指导填写。

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return raw(
        output.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="计费规则导入模板.xlsx"'},
    )


def _build_pricing_rules_excel(rules: list) -> bytes:
    """在独立线程中完成计费规则的 DataFrame 组装与 Excel 写入（纯 CPU）

    只访问 selectinload 预加载的 customer 关系与 JSON/日期普通列，不会在另一线程
    触发懒加载；供 export 端点经 ``asyncio.to_thread`` 调用，避免 5 万行规模下
    阻塞事件循环。
    """
    import pandas as pd

    # 组装 DataFrame（列与列表对齐）
    data = [
        {
            "id": r.id,
            # 与导入模板对齐：导入以 company_id（客户编号）为主键列。导出若给
            # customer_id（数据库内部主键），导出文件将无法回灌（报「缺少必填列：company_id」），
            # 与 balances 导出的 company_id / 「客户ID」约定也不一致。
            "company_id": r.customer.company_id if r.customer else None,
            "customer_name": r.customer.name if r.customer else None,
            "device_type": r.device_type,
            # 包年规则 layer_type 可为 NULL（模型注释「为空表示通用」）。导出用 or "single"
            # 会把 NULL 写成 "single"，回灌后规则 layer_type 从 NULL 被改成 "single"，
            # 往返不闭合；结算层已按 (device_type, layer_type or "single") 兜底，导出无需填默认值。
            "layer_type": r.layer_type,
            "pricing_type": r.pricing_type,
            # 真值判断会把 0 元单价/加层单价当成 None，须显式判 None
            "unit_price": float(r.unit_price) if r.unit_price is not None else None,
            "multi_floor_pricing_type": r.multi_floor_pricing_type,
            "additional_floor_price": (
                float(r.additional_floor_price) if r.additional_floor_price is not None else None
            ),
            "tiers": _json.dumps(r.tiers, ensure_ascii=False) if r.tiers else None,
            "package_type": r.package_type,
            # 必须按 CST 日期输出：DB 里存的是 UTC 时刻（CST 当日 00:00 -> UTC 前一日 16:00），
            # 直接 isoformat() 会得到 "2026-03-31T16:00:00+00:00"，导入端按 str()[:10] 取日期会
            # 得到 2026-03-31 并按 CST 重新解析 -> 导出再导入整体提前一天，往返不闭合。
            "effective_date": utc_to_cst_date_str(r.effective_date),
            "expiry_date": utc_to_cst_date_str(r.expiry_date),
        }
        for r in rules
    ]
    df = pd.DataFrame(data)

    output = io.BytesIO()
    df.to_excel(output, index=False, sheet_name="计费规则", engine="openpyxl")  # pyright: ignore[reportArgumentType]  # pandas-stubs 的 WriteExcelBuffer 未含 BytesIO（运行时支持）

    return output.getvalue()


@billing_bp.get("/pricing-rules/export")
@auth_required
@require_permission("billing:pricing_export")
async def export_pricing_rules(request: Request):
    """导出计费规则为 Excel（按当前筛选条件导出全部匹配数据，上限 50000 条）"""
    db: AsyncSession = request.ctx.db_session
    pricing_service = PricingService(PricingRepository(db))

    # 筛选参数（与列表页一致）
    try:
        customer_id = (
            int(request.args.get("customer_id")) if request.args.get("customer_id") else None
        )
    except ValueError:
        return json({"code": 40001, "message": "customer_id 参数必须为整数"}, status=400)
    keyword = request.args.get("keyword")
    device_type = request.args.get("device_type")
    layer_type = request.args.get("layer_type")
    pricing_type = request.args.get("pricing_type")

    rules, _total = await pricing_service.get_pricing_rules(
        customer_id=customer_id,
        keyword=keyword,
        device_type=device_type,
        layer_type=layer_type,
        pricing_type=pricing_type,
        page=1,
        page_size=50000,
    )

    if not rules:
        return json(
            {"code": 40002, "message": "没有找到符合条件的计费规则"},
            status=400,
        )

    # 行组装 + DataFrame + Excel 生成是纯 CPU 工作，5 万行规模下会阻塞事件循环，
    # 与 balances 导出一致移入线程（DB 查询留在异步层）；组装只访问 selectinload
    # 预加载的 customer 属性，不会在另一线程触发懒加载。
    excel_bytes = await asyncio.to_thread(_build_pricing_rules_excel, rules)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"pricing_rules_{timestamp}.xlsx"

    return raw(
        excel_bytes,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ==================== 结算单管理 ====================
