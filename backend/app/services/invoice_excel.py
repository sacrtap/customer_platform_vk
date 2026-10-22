"""结算单明细 Excel 生成服务

从外部数据库（nest_model_order）查询订单明细，生成包含
"合计"和"月份明细"两个 Sheet 的 Excel 文件。

模板参照: 安溪如是VR房源模型消费2026年7月结算单.xlsx
"""

import asyncio
import logging
import os
import re
from datetime import datetime, timedelta
from decimal import Decimal
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncEngine, AsyncSession

from ..config import settings
from ..utils.timezone import utc_to_cst_date_str

logger = logging.getLogger(__name__)

# Excel 单元格（XML 1.0）不允许的控制字符：openpyxl 写入时抛 IllegalCharacterError。
# 合法范围：\\t(0x09) \\n(0x0A) \\r(0x0D) 及 0x20 以上；剔除其余 C0 控制符。
_ILLEGAL_XML_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

# ============================================================
# 明细 Sheet 表头（29 列，与模板完全对齐）
# ============================================================
DETAIL_HEADERS = [
    "房源名称",
    "房源id",
    "城市",
    "业务部门",
    "订单编号",
    "项目链接",
    "创建人",
    "创建时间",
    "所属人",
    "面积",
    "订单状态",
    "编辑状态",
    "备注",
    "编辑人账号",
    "编辑人姓名",
    "上传完成时间",
    "处理完成时间",
    "第一次分配时间",
    "uv",
    "pv",
    "户型图绘制时间",
    "户型图最新绘制人",
    "审核发布时间",
    "楼层层数",
    "设备编号",
    "接单时间",
    "复审状态",
    "发布人",
    "户型图修模状态",
]

# ============================================================
# 订单状态映射
# ============================================================
ORDER_STATUS_MAP = {
    0: "拍摄中",
    2: "已接单未上传",
    3: "已上传处理中",
    5: "待编辑",
    6: "有编辑",
    8: "已上线",
    9: "有编辑已上线",
    11: "已下线",
    12: "有编辑已下线",
    13: "已取消",
    14: "拒绝VR拍摄",
    15: "已作废",
}

# 编辑状态映射（基于 produce_status + produce_flag）
EDIT_STATUS_MAP = {
    ("5", "1"): "待编辑",
    ("6", "1"): "已编辑户型图",
    ("6", "2"): "已编辑VR讲房",
    ("6", "3"): "已编辑户型图/已编辑VR讲房",
}

# 复审状态映射
REVIEW_STATUS_MAP = {
    0: "--",
    1: "待审核",
    2: "审核通过",
    3: "审核驳回",
}

# 户型图修模状态映射
MODEL_FIX_STATUS_MAP = {
    0: "--",
    1: "待修模",
    2: "修模中",
    3: "修模完成",
}

# ============================================================
# 明细数据查询 SQL
# 直接返回 datetime 对象（不使用 DATE_FORMAT），由 openpyxl 写入
# ============================================================
DETAIL_SQL = """
SELECT
    D.project_name,
    D.custom_code,
    D.city,
    D.department,
    D.order_code,
    D.nest_id,
    U_creator.owner_name AS creator_name,
    D.create_date,
    U_personal.owner_name AS personal_name,
    D.nest_area,
    D.order_status,
    D.produce_status,
    D.produce_flag,
    D.remarks,
    U_editor.username AS editor_account,
    U_editor.owner_name AS editor_name,
    D.upload_date,
    D.finish_date,
    D.first_send_zb_date,
    D.first_edit_time,
    D.review_release_date,
    D.floor_count,
    D.device_name,
    D.customer_review_result,
    U_publisher.owner_name AS publisher_name,
    D.model_fix_status
FROM nest_model_order D
LEFT JOIN nest_user U_creator ON U_creator.id = D.creator
LEFT JOIN nest_user U_personal ON U_personal.id = D.personal
LEFT JOIN nest_user U_editor ON U_editor.id = D.personal
LEFT JOIN nest_user U_publisher ON U_publisher.id = D.publisher
WHERE D.group_type = :group_type
  AND D.upload_date >= :start_dt AND D.upload_date < :end_dt
  AND ((D.order_status >= 3 AND D.order_status <= 12) OR D.order_status = 15)
ORDER BY D.create_date DESC
"""


class InvoiceExcelService:
    """结算单明细 Excel 生成服务"""

    def __init__(
        self,
        db: AsyncSession,
        external_engine: Optional[AsyncEngine] = None,
        session_factory: Optional[Any] = None,
    ):
        self.db = db
        self.external_engine = external_engine
        # 可选会话工厂：提供时余额查询改用短会话（查完即还），
        # 调用方可在构建 Excel 前提前关闭长事务会话（见
        # tasks/invoice_detail_generator.py 的事务缩短）。
        self.session_factory = session_factory

    async def generate_detail_file(
        self,
        invoice_id: int,
        customer_id: int,
        customer_name: str,
        period_start: datetime,
        period_end: datetime,
        total_amount: Decimal,
        discount_amount: Decimal,
        unit_price: Optional[Decimal] = None,
        group_type: Optional[int] = None,
        invoice_status: str = "draft",
        invoice_items: Optional[List[Dict[str, Any]]] = None,
    ) -> Optional[str]:
        """生成结算单明细 Excel 文件

        Returns:
            文件路径（成功）或 None（失败）
        """
        logger.info(f"开始生成结算单 {invoice_id} 的明细文件")

        # 1. 查询外部数据库获取订单明细
        order_details = await self._fetch_order_details(group_type, period_start, period_end)

        if order_details is None:
            logger.warning(f"结算单 {invoice_id} 查询外部数据库失败")
            return None

        logger.info(f"结算单 {invoice_id} 查询到 {len(order_details)} 条订单明细")

        # 2. 查询客户余额和当月充值
        balance_info = await self._fetch_balance_info(
            customer_id,
            period_start,
            period_end,
            total_amount,
            discount_amount,
            invoice_status,
        )

        # 3. 生成 Excel
        # _build_and_save_excel 是纯同步函数（openpyxl 逐格写样式 + wb.save zip
        # 序列化，零 await 点），其输入 order_details/balance_info 均为已取回的
        # 纯内存数据、函数内不碰 DB，线程安全。明细 job 跑在 AsyncIOScheduler
        # 的 AsyncIOExecutor 上（与 Web 服务同一事件循环），直接同步调用会独占
        # 循环数秒，导致批量生成期间列表/轮询/详情请求全部排队（实测详情 13ms
        # 被阻塞到 17s）。故丢默认线程池执行，事件循环保持响应。
        file_path = await asyncio.to_thread(
            self._build_and_save_excel,
            invoice_id=invoice_id,
            customer_name=customer_name,
            period_start=period_start,
            period_end=period_end,
            total_amount=total_amount,
            discount_amount=discount_amount,
            unit_price=unit_price,
            order_details=order_details,
            balance_info=balance_info,
            invoice_items=invoice_items or [],
        )

        logger.info(f"结算单 {invoice_id} 明细文件已生成: {file_path}")
        return file_path

    async def _fetch_order_details(
        self,
        group_type: Optional[int],
        period_start: datetime,
        period_end: datetime,
    ) -> Optional[List[Dict[str, Any]]]:
        """从外部 MySQL 查询订单明细

        Returns:
            订单列表（成功）或 None（连接失败）
        """
        if not group_type:
            logger.warning("group_type 未提供，无法查询外部数据库")
            return []

        # period_start/end 已经是 UTC datetime，直接使用
        start_dt = period_start
        end_dt = period_end + timedelta(seconds=1)  # 包含结束时刻

        if self.external_engine:
            try:
                async with self.external_engine.connect() as conn:
                    result = await conn.execute(
                        text(DETAIL_SQL),
                        {
                            "group_type": group_type,
                            "start_dt": start_dt,
                            "end_dt": end_dt,
                        },
                    )
                    rows = result.fetchall()
                    return [self._row_to_dict(row) for row in rows]
            except Exception as e:
                logger.error(f"查询外部数据库失败: {e}")
                return None
        else:
            logger.warning("外部 MySQL 引擎未配置，跳过明细查询")
            return []

    async def _fetch_balance_info(
        self,
        customer_id: int,
        period_start: datetime,
        period_end: datetime,
        total_amount: Decimal,
        discount_amount: Decimal,
        invoice_status: str,
    ) -> Dict[str, Any]:
        """查询客户余额和当月充值信息

        期初余额 = 当前余额 + 本结算单最终金额（如已扣款则需加回）
        当月充值 = 结算周期内充值记录总和
        结算后余额 = 期初余额 + 当月充值 - 最终结算金额
        """
        # 有会话工厂时用短会话查完即还：调用方（明细任务）在查询计费项后
        # 即关闭长事务会话，余额查询不应再依赖它
        if self.session_factory is not None:
            async with self.session_factory() as short_session:
                return await self._query_balance_info(
                    short_session,
                    customer_id,
                    period_start,
                    period_end,
                    total_amount,
                    discount_amount,
                    invoice_status,
                )

        return await self._query_balance_info(
            self.db,
            customer_id,
            period_start,
            period_end,
            total_amount,
            discount_amount,
            invoice_status,
        )

    async def _query_balance_info(
        self,
        db: AsyncSession,
        customer_id: int,
        period_start: datetime,
        period_end: datetime,
        total_amount: Decimal,
        discount_amount: Decimal,
        invoice_status: str,
    ) -> Dict[str, Any]:
        """余额/充值查询实现（db 为调用方提供的会话，由 _fetch_balance_info 分流）"""
        from sqlalchemy import func, select

        from ..models.billing import CustomerBalance, RechargeRecord

        result: Dict[str, Any] = {
            "opening_balance": Decimal(0),
            "monthly_recharge": Decimal(0),
            "closing_balance": Decimal(0),
        }

        try:
            # 查询当前余额
            balance = await db.execute(
                select(CustomerBalance).where(CustomerBalance.customer_id == customer_id)
            )
            bal = balance.scalars().first()

            current_balance = bal.total_amount if bal and bal.total_amount else Decimal(0)

            # 如果结算单已完成（已扣款），需要加回扣款金额恢复期初余额
            final_amount = total_amount - discount_amount
            if invoice_status in ("completed",):
                opening_balance = current_balance + final_amount
            else:
                opening_balance = current_balance

            # 查询当月充值记录总额
            # recharge_records.created_at 是 timestamp WITHOUT TIME ZONE（naive），
            # 而 period_start/end 来自 timestamptz 列的 aware UTC datetime。直接比较
            # asyncpg 抛 DataError（can't subtract offset-naive and offset-aware），
            # 该异常曾被 except 吞掉仅记 error，导致「当月充值」恒为 0。
            # 此处显式剥离 tzinfo（UTC 墙钟值不变）以匹配 naive 列语义。
            start_dt = period_start.replace(tzinfo=None)
            end_dt = (period_end + timedelta(seconds=1)).replace(tzinfo=None)  # 包含结束时刻
            recharge_result = await db.execute(
                select(
                    func.coalesce(func.sum(RechargeRecord.real_amount), 0),
                    func.coalesce(func.sum(RechargeRecord.bonus_amount), 0),
                ).where(
                    RechargeRecord.customer_id == customer_id,
                    RechargeRecord.created_at >= start_dt,
                    RechargeRecord.created_at < end_dt,
                )
            )
            recharge_row = recharge_result.fetchone()
            monthly_recharge = Decimal(0)
            if recharge_row:
                monthly_recharge = (recharge_row[0] or 0) + (recharge_row[1] or 0)

            # 结算后余额 = 期初 + 当月充值 - 最终结算金额
            closing_balance = opening_balance + monthly_recharge - final_amount

            result["opening_balance"] = opening_balance
            result["monthly_recharge"] = monthly_recharge
            result["closing_balance"] = closing_balance

        except Exception as e:
            logger.error(f"查询客户余额失败: {e}")

        return result

    def _row_to_dict(self, row) -> Dict[str, Any]:
        """将查询结果行转换为字典

        所有字符串字段经 _clean_cell_text 清洗：外部 MySQL 数据可能含控制字符
        （如项目名 "景发-怡景-12C\x0c"），openpyxl 写入时抛 IllegalCharacterError
        导致整个明细生成失败。datetime/数值字段原样返回。
        """
        raw = {
            "project_name": row[0] or "",
            "custom_code": row[1] or "",
            "city": row[2] or "",
            "department": row[3] or "",
            "order_code": row[4] or "",
            "nest_id": row[5] or "",
            "creator_name": row[6] or "",
            "create_date": row[7],  # datetime 对象
            "personal_name": row[8] or "",
            "nest_area": row[9] or "",
            "order_status": row[10],
            "produce_status": row[11],
            "produce_flag": row[12],
            "remarks": row[13] or "",
            "editor_account": row[14] or "",
            "editor_name": row[15] or "",
            "upload_date": row[16],  # datetime 对象
            "finish_date": row[17],  # datetime 对象
            "first_send_zb_date": row[18],  # datetime 对象
            "first_edit_time": row[19],  # datetime 对象
            "review_release_date": row[20],  # datetime 对象
            "floor_count": row[21] or "",
            "device_name": row[22] or "",
            "customer_review_result": row[23],
            "publisher_name": row[24] or "",
            "model_fix_status": row[25],
        }
        return {
            key: self._clean_cell_text(val) if isinstance(val, str) else val
            for key, val in raw.items()
        }

    @staticmethod
    def _clean_cell_text(val: Any) -> Any:
        """清洗单元格文本：移除 openpyxl 拒绝的控制字符（IllegalCharacterError）

        Excel 单元格合法字符为 XML 1.0 允许范围（含 \\t \\n \\r）；外部数据里的
        \\x00-\\x08 / \\x0b / \\x0c / \\x0e-\\x1f 等控制字符会让 openpyxl 在写入时
        抛 IllegalCharacterError，整个明细文件生成失败。此处统一剔除，其余原样保留。
        """
        if not isinstance(val, str):
            return val
        return _ILLEGAL_XML_CHARS_RE.sub("", val)

    @staticmethod
    def _fmt_dt(val: Any) -> str:
        """将 datetime 对象格式化为字符串 'YYYY-MM-DD HH:MM:SS'"""
        if val is None:
            return ""
        if isinstance(val, datetime):
            return val.strftime("%Y-%m-%d %H:%M:%S")
        return str(val)

    @staticmethod
    def _fmt_edit_status(produce_status: Any, produce_flag: Any) -> str:
        """推断编辑状态"""
        key = (str(produce_status or ""), str(produce_flag or ""))
        return EDIT_STATUS_MAP.get(key, "")

    @staticmethod
    def _fmt_review_status(val: Any) -> str:
        return REVIEW_STATUS_MAP.get(val, "--")

    @staticmethod
    def _fmt_model_fix_status(val: Any) -> str:
        return MODEL_FIX_STATUS_MAP.get(val, "--")

    def _build_and_save_excel(
        self,
        invoice_id: int,
        customer_name: str,
        period_start: datetime,
        period_end: datetime,
        total_amount: Decimal,
        discount_amount: Decimal,
        unit_price: Optional[Decimal],
        order_details: List[Dict[str, Any]],
        balance_info: Optional[Dict[str, Any]] = None,
        invoice_items: Optional[List[Dict[str, Any]]] = None,
    ) -> str:
        """生成 Excel 并保存到文件

        Returns:
            文件相对路径
        """
        wb = Workbook()

        # 通用样式
        header_font = Font(bold=True)
        header_fill = PatternFill("solid", start_color="D9E1F2")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # ============================================================
        # Sheet 1: 合计（10 列 A-J，与模板对齐）
        # ============================================================
        # types-openpyxl 将 cell(value) 声明为 str | None，但 openpyxl 运行时接受
        # int/float/Decimal；本函数大量写入数值列，标注 Any 统一豁免该误报。
        ws1: Any = wb.active
        assert ws1 is not None  # 新建 Workbook 必有活动工作表
        ws1.title = "合计"

        # 标题行
        ws1.merge_cells("A1:J1")
        ws1.cell(
            row=1,
            column=1,
            value=f"{self._clean_cell_text(customer_name)}VR房源模型消费结算汇总",
        )
        ws1.cell(row=1, column=1).font = Font(bold=True, size=14)
        ws1.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")

        # 表头（10 列，移除了"单价"列）
        headers_1 = [
            "结算周期",
            "模型总数量（套）",
            "计费模型数量（套）",
            "总金额（元）",
            "减免金额（元）",
            "实际结算金额（元）",
            "期初余额（元）",
            "当月充值金额（元）",
            "结算后余额（元）",
            "备注",
        ]
        for col, h in enumerate(headers_1, 1):
            cell = ws1.cell(row=2, column=col, value=h)
            cell.font = header_font
            cell.alignment = header_align
            cell.fill = header_fill

        # 数据行
        model_count = len(order_details)
        final_amount = total_amount - discount_amount
        period_str = f"{utc_to_cst_date_str(period_start)} - {utc_to_cst_date_str(period_end)}"
        ws1.cell(row=3, column=1, value=period_str)
        ws1.cell(row=3, column=2, value=model_count)
        ws1.cell(row=3, column=3, value=model_count)
        ws1.cell(row=3, column=4, value=float(total_amount))
        ws1.cell(row=3, column=5, value=float(discount_amount))
        ws1.cell(row=3, column=6, value=float(final_amount))
        # 期初余额 / 当月充值 / 结算后余额
        if balance_info:
            ws1.cell(row=3, column=7, value=float(balance_info.get("opening_balance", 0)))
            ws1.cell(row=3, column=8, value=float(balance_info.get("monthly_recharge", 0)))
            ws1.cell(
                row=3, column=9, value=float(balance_info.get("closing_balance", final_amount))
            )
        else:
            ws1.cell(row=3, column=7, value="")
            ws1.cell(row=3, column=8, value="")
            ws1.cell(row=3, column=9, value=float(final_amount))
        ws1.cell(row=3, column=10, value="")

        # 列宽
        for col in range(1, 11):
            ws1.column_dimensions[get_column_letter(col)].width = 18

        # ============================================================
        # Sheet 2: 月份明细（29 列，与模板对齐）
        # ============================================================
        ws2 = wb.create_sheet(title=f"{utc_to_cst_date_str(period_start)[:7]}")

        # 表头
        for col, h in enumerate(DETAIL_HEADERS, 1):
            cell = ws2.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.alignment = header_align
            cell.fill = header_fill

        # 数据行
        for row_idx, order in enumerate(order_details, 2):
            nest_id = order.get("nest_id", "")
            project_link = f"https://beyond.3dnest.cn/house/?m={nest_id}" if nest_id else ""
            status_val = order.get("order_status")
            status_str = (
                ORDER_STATUS_MAP.get(status_val, str(status_val)) if status_val is not None else ""
            )

            row_data = [
                order.get("project_name", ""),
                order.get("custom_code", ""),
                order.get("city", ""),
                order.get("department", ""),
                order.get("order_code", ""),
                project_link,
                order.get("creator_name", ""),
                self._fmt_dt(order.get("create_date")),
                order.get("personal_name", ""),
                order.get("nest_area", ""),
                status_str,
                self._fmt_edit_status(order.get("produce_status"), order.get("produce_flag")),
                order.get("remarks", ""),
                order.get("editor_account", ""),
                order.get("editor_name", ""),
                self._fmt_dt(order.get("upload_date")),
                self._fmt_dt(order.get("finish_date")),
                self._fmt_dt(order.get("first_send_zb_date")),
                0,  # uv
                0,  # pv
                self._fmt_dt(order.get("first_edit_time")),
                "",  # 户型图最新绘制人（需额外关联，暂留空）
                self._fmt_dt(order.get("review_release_date")),
                order.get("floor_count", ""),
                order.get("device_name", ""),
                "",  # 接单时间（需额外查询，暂留空）
                self._fmt_review_status(order.get("customer_review_result")),
                order.get("publisher_name", ""),
                self._fmt_model_fix_status(order.get("model_fix_status")),
            ]

            for col, val in enumerate(row_data, 1):
                # 兜底清洗：即使调用方传入未清洗的 order dict（含控制字符），
                # openpyxl 也不会抛 IllegalCharacterError。正常链路 _row_to_dict
                # 出口已清洗，这里防御其他构造来源。
                cell = ws2.cell(row=row_idx, column=col, value=self._clean_cell_text(val))
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        # 列宽
        col_widths = {
            1: 22,
            2: 12,
            3: 16,
            4: 12,
            5: 22,
            6: 50,
            7: 12,
            8: 22,
            9: 12,
            10: 10,
        }
        for col, w in col_widths.items():
            ws2.column_dimensions[get_column_letter(col)].width = w
        for col in range(11, len(DETAIL_HEADERS) + 1):
            if col not in col_widths:
                ws2.column_dimensions[get_column_letter(col)].width = 16

        # ============================================================
        # Sheet 3: 计费明细（InvoiceItem 列表）
        # ============================================================
        if invoice_items:
            ws3 = wb.create_sheet(title="计费明细")

            # 计费明细表头（7 列）
            billing_headers = [
                "设备类型",
                "楼层",
                "数量",
                "单价（元）",
                "小计（元）",
                "多层计费类型",
                "其他层单价（元）",
            ]
            for col, h in enumerate(billing_headers, 1):
                cell = ws3.cell(row=1, column=col, value=h)
                cell.font = header_font
                cell.alignment = header_align
                cell.fill = header_fill

            # 数据行
            for row_idx, item in enumerate(invoice_items, 2):
                # 设备类型映射
                dt = item.get("device_type") or ""
                dt_str = {"X": "X", "N": "N", "L": "L"}.get(dt, dt or "包年")

                # 楼层映射
                lt = item.get("layer_type", "")
                lt_str = {"single": "单层", "multi": "多层"}.get(lt, lt or "-")

                quantity = float(item.get("quantity", 0))
                unit_p = float(item.get("unit_price", 0))
                subtotal = quantity * unit_p

                mfp_type = item.get("multi_floor_pricing_type") or ""
                mfp_str = {"incremental": "递增", "unified": "统一"}.get(mfp_type, mfp_type or "-")

                afp = item.get("additional_floor_price")
                afp_val = float(afp) if afp is not None else ""

                row_data = [dt_str, lt_str, quantity, unit_p, subtotal, mfp_str, afp_val]
                for col, val in enumerate(row_data, 1):
                    cell = ws3.cell(row=row_idx, column=col, value=val)
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center")

            # 列宽
            billing_widths = {1: 12, 2: 10, 3: 10, 4: 12, 5: 12, 6: 14, 7: 16}
            for col, w in billing_widths.items():
                ws3.column_dimensions[get_column_letter(col)].width = w

        # ============================================================
        # 保存文件
        # ============================================================
        year = utc_to_cst_date_str(period_start)[:4]
        month = utc_to_cst_date_str(period_start)[5:7]
        rel_dir = os.path.join("invoices", year, month)
        base_dir = getattr(settings, "file_storage_path", "./uploads")
        abs_dir = os.path.join(base_dir, rel_dir)
        os.makedirs(abs_dir, exist_ok=True)

        filename = f"invoice_{invoice_id}.xlsx"
        abs_path = os.path.join(abs_dir, filename)
        wb.save(abs_path)

        # 落盘后二次确认：文件必须存在且非空，状态才可置 completed
        # 若文件写入异常（磁盘满、权限等），抛出异常让调用方置 failed
        if not os.path.exists(abs_path) or os.path.getsize(abs_path) == 0:
            raise IOError(f"明细文件写入后验证失败（不存在或为空）: {abs_path}")

        return os.path.join(rel_dir, filename)
